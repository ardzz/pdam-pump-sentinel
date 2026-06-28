from pathlib import Path
from zipfile import ZipFile

from pptx import Presentation

from scripts.generate_presentation_deck import CANONICAL_SCREENSHOTS, OUTPUT_PATH, SLIDES, build_deck

ROOT = Path(__file__).resolve().parents[2]
OUTLINE = ROOT / 'docs/presentation/deck-outline.md'


def test_deck_outline_has_expected_slide_contract():
    text = OUTLINE.read_text(encoding='utf-8')

    assert text.count('## Slide ') == 13
    assert 'SKAB dipakai sebagai surrogate water-circulation testbed' in text
    assert 'local Docker Compose evidence' in text
    for screenshot in CANONICAL_SCREENSHOTS:
        assert screenshot in text
        assert (ROOT / 'docs/presentation/screenshots' / screenshot).exists()


def test_generated_pptx_is_structurally_valid(tmp_path):
    output = tmp_path / 'deck.pptx'

    build_deck(output)

    assert output.exists()
    with ZipFile(output) as archive:
        assert archive.testzip() is None

    presentation = Presentation(str(output))
    assert len(presentation.slides) == len(SLIDES) == 13

    text = '\n'.join(
        _shape_text(shape)
        for slide in presentation.slides
        for shape in slide.shapes
    ).lower()
    forbidden_fragments = ('api key', 'password=', 'secret=', 'real pdam production data')
    for fragment in forbidden_fragments:
        assert fragment not in text


def test_committed_deck_artifact_exists_after_generation():
    if not OUTPUT_PATH.exists():
        build_deck(OUTPUT_PATH)

    presentation = Presentation(str(OUTPUT_PATH))

    assert len(presentation.slides) == 13


def _shape_text(shape: object) -> str:
    text = getattr(shape, 'text', '')
    return text if isinstance(text, str) else ''
